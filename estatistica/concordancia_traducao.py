# concordancia_traducao.py
# Autor: Guilherme Lacerda de Avila
#
# Script pra calcular a PORCENTAGEM DE CONCORDÂNCIA entre as versões
# traduzidas do RPMS - pedido da orientadora (18/08/2026) como método
# oficial pra análise da tradução (a retrotradução tem script próprio,
# ver concordancia_retrotraducao.py).
#
# IMPORTANTE - isso é diferente de tudo nos outros scripts: o
# Kappa/CVI/AC1/Alpha medem concordância entre JUÍZES avaliando ITENS
# (resposta binária Sim/Não). Aqui a gente mede o quão PARECIDO é o TEXTO
# entre versões traduzidas diferentes do mesmo item - não tem categoria
# Sim/Não nenhuma, é comparação de string.
#
# Fonte dos dados: "Versão Corrigida dos Tradutores.xlsx" (a versão mais
# recente/corrigida do processo de tradução - tem as 3 versões lado a
# lado, já reconciliadas pela equipe):
#   coluna C = Tradução (Português) - tradução inicial, leiga
#   coluna D = Tradução (Português - especialista)
#   coluna E = Versão Proposta Equipe - síntese final da equipe
#
# Repare que a coluna E normalmente NASCE de uma discussão em cima de C e
# D (não é uma tradução independente) - por isso o resultado mais
# "limpo" pra comparar dois tradutores de verdade é C vs D; C-vs-E e
# D-vs-E mostram o quanto a versão final se afastou de cada tradução
# inicial.
#
# O que conta como "concordância" aqui, calculado de dois jeitos:
#   1) EXATA - o texto das duas versões é idêntico depois de normalizar
#      (ignorando maiúscula/minúscula, espaços extras, numeração do item
#      tipo "6." no começo). É rígido: qualquer troca de palavra já conta
#      como diferente.
#   2) SIMILARIDADE - um percentual "de quanto o texto se parece", usando
#      difflib.SequenceMatcher (biblioteca padrão do Python, não precisa
#      instalar nada) - da 100% pra texto idêntico, 0% pra texto
#      totalmente diferente, e vai suavemente entre os dois pra parecido
#      "mas não igual" (o caso mais comum aqui).
#
# Como rodar:
#   uv run estatistica/concordancia_traducao.py
#   uv run estatistica/concordancia_traducao.py "caminho/outro_arquivo.xlsx"

import re
import sys
import textwrap
from datetime import datetime
from difflib import SequenceMatcher
from pathlib import Path

import matplotlib
import openpyxl

matplotlib.use("Agg")  # backend sem tela, so pra salvar arquivo de imagem direto
import matplotlib.pyplot as plt  # noqa: E402 (precisa vir depois do matplotlib.use)

ABA = "RPMS - Tradução e Comparações"

LINHA_TITULO = 2
LINHA_INSTRUCAO = 3
LINHA_ESCALA_INICIO = 4
LINHA_ESCALA_FIM = 8
LINHA_ITENS_INICIO = 9
LINHA_ITENS_FIM = 29

COLUNA_LEIGA = 3  # C - Tradução (Português)
COLUNA_ESPECIALISTA = 4  # D - Tradução (Português - especialista)
COLUNA_EQUIPE = 5  # E - Versão Proposta Equipe

PASTA_DO_SCRIPT = Path(__file__).resolve().parent
ARQUIVO_PADRAO = PASTA_DO_SCRIPT.parent / "Versão Corrigida dos Tradutores.xlsx"
PASTA_RESULTADOS = PASTA_DO_SCRIPT / "resultados"


def normalizar_texto(texto):
    """
    Deixa o texto pronto pra comparar: tira espaço em branco "invisível"
    (\\xa0, que apareceu em alguns itens da planilha), tira espaço extra
    no começo/fim, deixa tudo minúsculo, e tira a numeração do item que
    fica no começo do texto (tipo "6. " ou "8.a "), porque essa numeração
    é só rótulo - não é parte da tradução em si.
    """
    if texto is None:
        return ""

    texto = str(texto).replace("\xa0", " ").strip().lower()
    texto = re.sub(r"^\d+\s*\.?\s*[a-z]?\s*\.?\s*", "", texto)
    texto = re.sub(r"\s+", " ", texto).strip()

    return texto


def calcular_similaridade(texto_a, texto_b):
    # SequenceMatcher.ratio() da um numero de 0 a 1 (0 = nada parecido, 1
    # = identico) baseado em quantos "pedacos" de caractere batem entre
    # as duas strings - multiplico por 100 pra virar porcentagem
    return SequenceMatcher(None, texto_a, texto_b).ratio() * 100


def rotular_linha(linha):
    if linha == LINHA_TITULO:
        return "Título da escala"
    if linha == LINHA_INSTRUCAO:
        return "Instrução"
    if LINHA_ESCALA_INICIO <= linha <= LINHA_ESCALA_FIM:
        return "Opção de resposta " + str(linha - LINHA_ESCALA_INICIO + 1)
    numero_do_item = linha - LINHA_ITENS_INICIO + 1
    return "Item " + str(numero_do_item)


def carregar_elementos(caminho_excel):
    """
    Le a planilha e monta uma lista de "elementos comparaveis" - cada um
    e um dicionario com o rotulo (pra aparecer no relatorio) e o texto
    das 3 versoes traduzidas. Da linha 2 ate a 29: titulo, instrucao, as
    5 opcoes de resposta da escala, e os 21 itens.
    """
    livro = openpyxl.load_workbook(caminho_excel, data_only=True)
    planilha = livro[ABA]

    elementos = []
    for linha in range(LINHA_TITULO, LINHA_ITENS_FIM + 1):
        leiga = planilha.cell(row=linha, column=COLUNA_LEIGA).value
        especialista = planilha.cell(row=linha, column=COLUNA_ESPECIALISTA).value
        equipe = planilha.cell(row=linha, column=COLUNA_EQUIPE).value

        # algumas linhas podem estar vazias na planilha (linha em branco
        # de separacao) - pulo essas
        if leiga is None and especialista is None and equipe is None:
            continue

        elementos.append(
            {
                "rotulo": rotular_linha(linha),
                "leiga": leiga if leiga is not None else "",
                "especialista": especialista if especialista is not None else "",
                "equipe": equipe if equipe is not None else "",
            }
        )

    return elementos


def calcular_comparacoes(elementos):
    """
    Pra cada elemento, calcula concordancia exata e similaridade nos 3
    pares possiveis (leiga-especialista, leiga-equipe,
    especialista-equipe), alem de saber se os 3 juntos sao identicos.
    """
    resultados = []

    for elemento in elementos:
        leiga_norm = normalizar_texto(elemento["leiga"])
        especialista_norm = normalizar_texto(elemento["especialista"])
        equipe_norm = normalizar_texto(elemento["equipe"])

        exato_le = leiga_norm == especialista_norm
        exato_lq = leiga_norm == equipe_norm
        exato_eq = especialista_norm == equipe_norm
        exato_todos = leiga_norm == especialista_norm == equipe_norm

        sim_le = calcular_similaridade(leiga_norm, especialista_norm)
        sim_lq = calcular_similaridade(leiga_norm, equipe_norm)
        sim_eq = calcular_similaridade(especialista_norm, equipe_norm)

        resultados.append(
            {
                "rotulo": elemento["rotulo"],
                "exato_leiga_especialista": exato_le,
                "exato_leiga_equipe": exato_lq,
                "exato_especialista_equipe": exato_eq,
                "exato_todos": exato_todos,
                "sim_leiga_especialista": sim_le,
                "sim_leiga_equipe": sim_lq,
                "sim_especialista_equipe": sim_eq,
                "sim_media": (sim_le + sim_lq + sim_eq) / 3,
            }
        )

    return resultados


def calcular_agregados(resultados):
    total = len(resultados)

    return {
        "total_elementos": total,
        "pct_exato_leiga_especialista": 100 * sum(r["exato_leiga_especialista"] for r in resultados) / total,
        "pct_exato_leiga_equipe": 100 * sum(r["exato_leiga_equipe"] for r in resultados) / total,
        "pct_exato_especialista_equipe": 100 * sum(r["exato_especialista_equipe"] for r in resultados) / total,
        "pct_exato_todos": 100 * sum(r["exato_todos"] for r in resultados) / total,
        "sim_media_leiga_especialista": sum(r["sim_leiga_especialista"] for r in resultados) / total,
        "sim_media_leiga_equipe": sum(r["sim_leiga_equipe"] for r in resultados) / total,
        "sim_media_especialista_equipe": sum(r["sim_especialista_equipe"] for r in resultados) / total,
        "sim_media_geral": sum(r["sim_media"] for r in resultados) / total,
    }


def montar_texto_do_relatorio(resultados, agregados):
    linhas = []
    linhas.append("=" * 70)
    linhas.append("RESULTADO - CONCORDÂNCIA ENTRE VERSÕES TRADUZIDAS - RPMS")
    linhas.append("=" * 70)
    linhas.append("Fonte: Versão Corrigida dos Tradutores.xlsx")
    linhas.append("Comparando: Tradução leiga (C) x Tradução especialista (D) x Versão da equipe (E)")
    linhas.append("Total de elementos comparados: " + str(agregados["total_elementos"]))
    linhas.append("  (título + instrução + 5 opções de resposta + 21 itens da escala)")
    linhas.append("")

    linhas.append("--- Por elemento ---")
    for r in resultados:
        linhas.append(r["rotulo"] + ":")
        linhas.append(
            "  leiga×especialista: "
            + ("idêntico" if r["exato_leiga_especialista"] else "diferente")
            + " (similaridade "
            + str(round(r["sim_leiga_especialista"], 1))
            + "%)"
        )
        linhas.append(
            "  leiga×equipe: "
            + ("idêntico" if r["exato_leiga_equipe"] else "diferente")
            + " (similaridade "
            + str(round(r["sim_leiga_equipe"], 1))
            + "%)"
        )
        linhas.append(
            "  especialista×equipe: "
            + ("idêntico" if r["exato_especialista_equipe"] else "diferente")
            + " (similaridade "
            + str(round(r["sim_especialista_equipe"], 1))
            + "%)"
        )

    linhas.append("")
    linhas.append("--- Resumo geral ---")
    linhas.append(
        "Concordância EXATA (as 3 versões idênticas): "
        + str(round(agregados["pct_exato_todos"], 1))
        + "% dos elementos"
    )
    linhas.append(
        "Concordância EXATA leiga×especialista: "
        + str(round(agregados["pct_exato_leiga_especialista"], 1))
        + "%"
    )
    linhas.append(
        "Concordância EXATA leiga×equipe: " + str(round(agregados["pct_exato_leiga_equipe"], 1)) + "%"
    )
    linhas.append(
        "Concordância EXATA especialista×equipe: "
        + str(round(agregados["pct_exato_especialista_equipe"], 1))
        + "%"
    )
    linhas.append("")
    linhas.append(
        "Similaridade MÉDIA leiga×especialista: "
        + str(round(agregados["sim_media_leiga_especialista"], 1))
        + "%"
    )
    linhas.append(
        "Similaridade MÉDIA leiga×equipe: " + str(round(agregados["sim_media_leiga_equipe"], 1)) + "%"
    )
    linhas.append(
        "Similaridade MÉDIA especialista×equipe: "
        + str(round(agregados["sim_media_especialista_equipe"], 1))
        + "%"
    )
    linhas.append("Similaridade MÉDIA GERAL (os 3 pares juntos): " + str(round(agregados["sim_media_geral"], 1)) + "%")
    linhas.append("")

    piores = sorted(resultados, key=lambda r: r["sim_media"])[:5]
    linhas.append("--- 5 elementos com MAIOR divergência entre as versões (menor similaridade média) ---")
    for r in piores:
        linhas.append("  " + r["rotulo"] + ": " + str(round(r["sim_media"], 1)) + "% de similaridade média")

    linhas.append("=" * 70)

    return "\n".join(linhas)


def montar_resumo_para_figura(agregados):
    colunas = [
        "Elementos (N)",
        "% exato (3 versões)",
        "% exato leiga×espec.",
        "Similaridade média geral",
    ]
    valores = [
        str(agregados["total_elementos"]),
        str(round(agregados["pct_exato_todos"], 1)) + "%",
        str(round(agregados["pct_exato_leiga_especialista"], 1)) + "%",
        str(round(agregados["sim_media_geral"], 1)) + "%",
    ]

    titulo = "Concordância entre versões traduzidas: RPMS (leiga × especialista × equipe)"

    notas = [
        "Nota. \"% exato\" = as versões comparadas são idênticas depois de normalizar (minúsculas, "
        "espaços, numeração do item removida) - medida rígida, qualquer troca de palavra já conta "
        "como diferente.",
        "\"Similaridade\" = percentual calculado com difflib.SequenceMatcher (Python), mede o quanto "
        "o texto se parece mesmo quando não é idêntico - mais informativo que o % exato pra tradução "
        "livre, onde reformular com outras palavras é esperado e não é erro.",
    ]

    return {"titulo": titulo, "colunas": colunas, "valores": valores, "notas": notas}


def salvar_relatorio_em_arquivo(texto_relatorio, agora):
    PASTA_RESULTADOS.mkdir(parents=True, exist_ok=True)

    nome_do_arquivo = "concordancia_traducao_" + agora + ".txt"
    caminho_completo = PASTA_RESULTADOS / nome_do_arquivo

    caminho_completo.write_text(texto_relatorio, encoding="utf-8")

    return caminho_completo


def salvar_imagem_relatorio(resumo, agora):
    PASTA_RESULTADOS.mkdir(parents=True, exist_ok=True)

    plt.rcParams["font.family"] = "serif"

    titulo = resumo["titulo"]
    colunas = resumo["colunas"]
    valores = resumo["valores"]
    notas = resumo["notas"]

    linhas_notas = []
    for nota in notas:
        linhas_notas.extend(textwrap.wrap(nota, width=110))
    texto_notas = "\n".join(linhas_notas)

    altura_titulo = 0.5
    altura_tabela = 1.3
    altura_notas = 0.15 + 0.24 * max(len(linhas_notas), 1)
    altura_total = altura_titulo + altura_tabela + altura_notas

    fig = plt.figure(figsize=(10.5, altura_total))
    fig.patch.set_facecolor("white")

    fig.text(
        0.03,
        1 - (0.32 / altura_total),
        titulo,
        fontsize=11,
        style="italic",
        ha="left",
        va="top",
        wrap=True,
    )

    topo_tabela = 1 - (altura_titulo / altura_total)
    altura_tabela_frac = altura_tabela / altura_total
    ax = fig.add_axes([0.03, topo_tabela - altura_tabela_frac, 0.94, altura_tabela_frac])
    ax.axis("off")

    tabela = ax.table(cellText=[valores], colLabels=colunas, cellLoc="center", loc="center")
    tabela.auto_set_font_size(False)
    tabela.set_fontsize(10)
    tabela.auto_set_column_width(col=list(range(len(colunas))))
    tabela.scale(1, 2.2)

    for (linha, _coluna), celula in tabela.get_celld().items():
        celula.set_edgecolor("black")
        if linha == 0:
            celula.set_text_props(weight="bold")
            celula.visible_edges = "BT"
            celula.set_linewidth(1.1)
        else:
            celula.visible_edges = "B"
            celula.set_linewidth(1.1)

    y_notas = topo_tabela - altura_tabela_frac - (0.05 / altura_total)
    fig.text(0.03, y_notas, texto_notas, fontsize=8, ha="left", va="top")

    nome_do_arquivo = "concordancia_traducao_" + agora + ".png"
    caminho_completo = PASTA_RESULTADOS / nome_do_arquivo

    fig.savefig(caminho_completo, dpi=200, facecolor="white")
    plt.close(fig)

    return caminho_completo


def salvar_html_relatorio(resumo, nome_da_imagem, agora):
    PASTA_RESULTADOS.mkdir(parents=True, exist_ok=True)

    titulo = resumo["titulo"]
    colunas = resumo["colunas"]
    valores = resumo["valores"]
    notas = resumo["notas"]

    celulas_cabecalho = "".join("<th>" + c + "</th>" for c in colunas)
    celulas_dados = "".join("<td>" + v + "</td>" for v in valores)
    paragrafos_notas = "".join("<p>" + n + "</p>" for n in notas)

    html = (
        "<!doctype html>\n"
        '<html lang="pt-BR">\n'
        "<head>\n"
        '<meta charset="utf-8">\n'
        "<title>" + titulo + "</title>\n"
        "<style>\n"
        "  body { margin: 0; padding: 24px; background: #ffffff;\n"
        '    font-family: "Times New Roman", Times, Georgia, serif; color: #000000; }\n'
        "  .titulo-tabela { font-size: 14px; margin-bottom: 14px; font-style: italic; }\n"
        "  table { width: 100%; border-collapse: collapse; font-size: 13px; }\n"
        "  thead tr { border-top: 1.6px solid #000; }\n"
        "  th, td { padding: 7px 10px; text-align: center; }\n"
        "  thead th { border-bottom: 1px solid #000; font-weight: bold; }\n"
        "  tbody tr:last-child td { border-bottom: 1.6px solid #000; }\n"
        "  .nota { font-size: 11.5px; margin-top: 12px; line-height: 1.5; }\n"
        "  .nota p { margin: 4px 0; }\n"
        "  .imagem-recente { margin-top: 28px; }\n"
        "  .imagem-recente p { font-size: 11px; font-style: italic; margin-bottom: 6px; }\n"
        "  .imagem-recente img { max-width: 100%; border: 1px solid #ccc; }\n"
        "</style>\n"
        "</head>\n"
        "<body>\n"
        '  <div class="titulo-tabela">' + titulo + "</div>\n"
        "  <table>\n"
        "    <thead><tr>" + celulas_cabecalho + "</tr></thead>\n"
        "    <tbody><tr>" + celulas_dados + "</tr></tbody>\n"
        "  </table>\n"
        '  <div class="nota">' + paragrafos_notas + "</div>\n"
        '  <div class="imagem-recente">\n'
        "    <p>Imagem gerada nesta execução (" + agora + "):</p>\n"
        '    <img src="' + nome_da_imagem + '" alt="' + titulo + '">\n'
        "  </div>\n"
        "</body>\n"
        "</html>\n"
    )

    nome_do_arquivo = "concordancia_traducao_" + agora + ".html"
    caminho_completo = PASTA_RESULTADOS / nome_do_arquivo
    caminho_completo.write_text(html, encoding="utf-8")

    return caminho_completo


def main():
    if len(sys.argv) > 1:
        caminho_excel = Path(sys.argv[1])
    else:
        caminho_excel = ARQUIVO_PADRAO

    if not caminho_excel.exists():
        print("Não achei esse arquivo aqui: " + str(caminho_excel))
        sys.exit(1)

    elementos = carregar_elementos(caminho_excel)
    resultados = calcular_comparacoes(elementos)
    agregados = calcular_agregados(resultados)

    texto_relatorio = montar_texto_do_relatorio(resultados, agregados)
    print(texto_relatorio)

    agora = datetime.now().strftime("%Y%m%d_%H%M%S")

    caminho_salvo = salvar_relatorio_em_arquivo(texto_relatorio, agora)
    print("\nRelatório salvo em: " + str(caminho_salvo))

    resumo = montar_resumo_para_figura(agregados)

    caminho_imagem = salvar_imagem_relatorio(resumo, agora)
    print("Imagem (tabela-resumo pronta pro artigo) salva em: " + str(caminho_imagem))

    caminho_html = salvar_html_relatorio(resumo, caminho_imagem.name, agora)
    print("Página HTML (com a imagem mais recente) salva em: " + str(caminho_html))


if __name__ == "__main__":
    main()
