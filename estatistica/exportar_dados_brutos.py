# exportar_dados_brutos.py
# Autor: Guilherme Lacerda de Avila
#
# Gera os arquivos de dados brutos anonimizados que sao publicados no
# repositorio (dados/avaliacao_epmr.csv e dados/avaliacao_rpms.csv), a
# partir das planilhas Excel originais.
#
# POR QUE ESSE SCRIPT EXISTE
#
# As planilhas Excel originais NAO podem ser publicadas: o nome de cada
# aba tem o nome real do juiz, e da linha 28 em diante tem nome completo,
# idade e pais de origem. No caso do RPMS os avaliadores sao refugiados,
# entao esse dado e especialmente sensivel.
#
# Esse script le as planilhas, joga fora tudo que e informacao pessoal e
# escreve um CSV com o que interessa pra reproduzir os calculos: qual
# juiz (numerado, sem nome), qual item, o que ele respondeu, e a ressalva
# que ele escreveu, quando escreveu.
#
# QUEM PRECISA RODAR ISSO
#
# So quem tem as planilhas originais, ou seja, a equipe de pesquisa. Quem
# clonou o repositorio ja recebe os CSVs prontos em dados/ e roda os
# scripts de calculo direto, sem precisar deste aqui.
#
# Como rodar:
#   uv run estatistica/exportar_dados_brutos.py

import csv
import sys
from pathlib import Path

import openpyxl

LINHA_COMECO = 6
LINHA_FIM = 26
COLUNA_RESPOSTA = 3  # coluna C
COLUNA_JUSTIFICATIVA = 4  # coluna D

PASTA_DO_SCRIPT = Path(__file__).resolve().parent
PASTA_RAIZ = PASTA_DO_SCRIPT.parent
PASTA_DADOS = PASTA_RAIZ / "dados"

PLANILHAS = [
    ("Avaliação EPMR Juízes Especialistas.xlsx", "avaliacao_epmr.csv"),
    ("Avaliação RPMS - Juízes Não Especialistas.xlsx", "avaliacao_rpms.csv"),
]


def normalizar_resposta(valor):
    if valor is None:
        raise ValueError("achei uma celula vazia onde devia ter Sim ou Nao")

    texto = str(valor).strip().lower()

    if texto == "sim":
        return "Sim"
    elif texto == "não" or texto == "nao":
        return "Não"
    else:
        raise ValueError("valor estranho na planilha, não é Sim nem Não: " + repr(valor))


def limpar_texto_da_ressalva(valor):
    # tira quebra de linha e espaco sobrando, pra caber numa celula de CSV
    if valor is None:
        return ""
    texto = str(valor).replace("\n", " ").replace("\r", " ").strip()
    while "  " in texto:
        texto = texto.replace("  ", " ")
    return texto


def extrair(caminho_excel):
    livro = openpyxl.load_workbook(caminho_excel, data_only=True)

    linhas = []
    # o nome da aba tem nome de pessoa de verdade: ele so serve pra abrir a
    # aba certa aqui, e vira um numero sequencial na saida
    for numero_do_juiz, nome_da_planilha in enumerate(livro.sheetnames, start=1):
        planilha = livro[nome_da_planilha]

        for linha in range(LINHA_COMECO, LINHA_FIM + 1):
            numero_do_item = linha - LINHA_COMECO + 1

            resposta = normalizar_resposta(planilha.cell(row=linha, column=COLUNA_RESPOSTA).value)
            ressalva = limpar_texto_da_ressalva(
                planilha.cell(row=linha, column=COLUNA_JUSTIFICATIVA).value
            )

            linhas.append(
                {
                    "item": numero_do_item,
                    "juiz": numero_do_juiz,
                    "resposta": resposta,
                    "tem_ressalva": 1 if ressalva else 0,
                    "ressalva": ressalva,
                }
            )

    return linhas


def salvar_csv(linhas, caminho_saida):
    caminho_saida.parent.mkdir(parents=True, exist_ok=True)

    with caminho_saida.open("w", encoding="utf-8", newline="") as arquivo:
        escritor = csv.DictWriter(
            arquivo, fieldnames=["item", "juiz", "resposta", "tem_ressalva", "ressalva"]
        )
        escritor.writeheader()
        for linha in linhas:
            escritor.writerow(linha)


def main():
    houve_erro = False

    for nome_excel, nome_csv in PLANILHAS:
        caminho_excel = PASTA_RAIZ / nome_excel

        if not caminho_excel.exists():
            print("Não achei a planilha: " + str(caminho_excel))
            print("  (esse script só roda pra quem tem as planilhas originais)")
            houve_erro = True
            continue

        linhas = extrair(caminho_excel)
        caminho_saida = PASTA_DADOS / nome_csv
        salvar_csv(linhas, caminho_saida)

        total_de_juizes = len({linha["juiz"] for linha in linhas})
        total_de_itens = len({linha["item"] for linha in linhas})
        total_com_ressalva = sum(linha["tem_ressalva"] for linha in linhas)

        print("Gerado: " + str(caminho_saida))
        print(
            "  "
            + str(total_de_itens)
            + " itens × "
            + str(total_de_juizes)
            + " juízes = "
            + str(len(linhas))
            + " respostas, "
            + str(total_com_ressalva)
            + " com ressalva escrita"
        )

    if houve_erro:
        sys.exit(1)


if __name__ == "__main__":
    main()
