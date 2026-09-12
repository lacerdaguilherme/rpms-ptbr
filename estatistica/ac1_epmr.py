# ac1_epmr.py
# Autor: Guilherme Lacerda de Avila
#
# Script pra calcular SÓ o AC1 de Gwet do EPMR - sem CVI junto (isso já
# existe em cvi_ac1_epmr.py, que calcula os dois lado a lado). Esse script
# aqui existe porque a orientadora pediu explicitamente um produto com o
# AC1 isolado, já que é o único coeficiente que vai pro corpo do artigo
# (ver RESULTADOS_PARA_ARTIGO.md, seção 0).
#
# O que é o AC1 de Gwet e por que ele (e não o Kappa de Fleiss) - ver
# cvi_ac1_epmr.py e EXPLICACAO.md, seção 9, pra explicação completa. Resumo:
# é uma estatística "tipo kappa" (corrigida por concordância ao acaso), mas
# com uma fórmula pra concordância esperada ao acaso que NÃO sofre do
# paradoxo do kappa quando a prevalência das respostas é desbalanceada
# (quase tudo "Sim", pouco "Não" - exatamente o nosso caso).
#   Pa = concordância observada média
#   pi = proporção média de "Sim" entre TODAS as respostas
#   Pe_ac1 = 2 * pi * (1 - pi)
#   AC1 = (Pa - Pe_ac1) / (1 - Pe_ac1)
#
# Como rodar (o projeto usa uv pra gerenciar as dependencias):
#   uv run estatistica/ac1_epmr.py
#
# Se quiser usar outro arquivo excel (tipo uma versao corrigida da
# planilha), da pra passar o caminho dele na hora de rodar:
#   uv run estatistica/ac1_epmr.py "caminho/outro_arquivo.xlsx"
# Se nao passar nada ele usa o arquivo padrao que ja vem no projeto.
#
# REGRA DE INTERPRETACAO "Sim + sugestao" (revisada em 11/09/2026, apos
# auditoria celula a celula das planilhas): discordancia e o "Nao" marcado
# na coluna binaria (coluna C). Uma ressalva escrita na coluna D so e
# reclassificada como discordancia quando aponta defeito no item como
# escrito, e nao sugestao de aprimoramento de um item ja julgado coerente.
# Os itens reclassificados ficam listados explicitamente na constante
# ITENS_COM_RESSALVA_RECLASSIFICADA, abaixo, um por instrumento.
#
# A versao anterior deste script reclassificava QUALQUER celula preenchida
# na coluna D, o que derrubou o AC1 do EPMR de 0,895 para 0,735 e divergiu
# do numero apresentado no poster e do que a orientadora de fato aplicou
# (um unico caso, o item 6 do RPMS). Ver EXPLICACAO.md, secao 16.

import csv
import math
import sys
import textwrap
from datetime import datetime
from pathlib import Path

import matplotlib
import openpyxl

matplotlib.use("Agg")  # backend sem tela, so pra salvar arquivo de imagem direto
import matplotlib.pyplot as plt  # noqa: E402 (precisa vir depois do matplotlib.use)

LINHA_COMECO = 6
LINHA_FIM = 26  # da 21 itens no total (26 - 6 + 1)
COLUNA_RESPOSTA = 3  # coluna C
COLUNA_JUSTIFICATIVA = 4  # coluna D

# Itens cuja ressalva foi reclassificada como discordancia (ver cabecalho).
# No painel de juizes especialistas: NENHUM. As observacoes registradas nos
# itens 4, 5, 8, 10 e 17 acompanham itens marcados como coerentes e propoem
# aprimoramento de redacao ou acrescimo de conteudo, nao apontam defeito no
# item como escrito. As reprovacoes deste painel sao os "Nao" da coluna C
# (itens 13 e 15 pelo Juiz 1; itens 20 e 21 pelo Juiz 4).
ITENS_COM_RESSALVA_RECLASSIFICADA = frozenset()

PASTA_DO_SCRIPT = Path(__file__).resolve().parent
ARQUIVO_PADRAO = PASTA_DO_SCRIPT.parent / "dados" / "avaliacao_epmr.csv"
PASTA_RESULTADOS = PASTA_DO_SCRIPT / "resultados"


def normalizar_resposta(valor):
    if valor is None:
        raise ValueError("achei uma celula vazia onde devia ter Sim ou Nao")

    texto = str(valor).strip().lower()

    if texto == "sim":
        return "Sim"
    elif texto == "não" or texto == "nao":
        return "Não"
    else:
        raise ValueError("valor estranho na planilha, não é Sim nem Não: " + repr(valor))


def ler_respostas_de_um_juiz(planilha):
    respostas = []
    for linha in range(LINHA_COMECO, LINHA_FIM + 1):
        valor_da_celula = planilha.cell(row=linha, column=COLUNA_RESPOSTA).value
        resposta = normalizar_resposta(valor_da_celula)

        numero_do_item = linha - LINHA_COMECO + 1
        valor_da_justificativa = planilha.cell(row=linha, column=COLUNA_JUSTIFICATIVA).value
        tem_ressalva = valor_da_justificativa is not None and str(valor_da_justificativa).strip() != ""
        if resposta == "Sim" and tem_ressalva and numero_do_item in ITENS_COM_RESSALVA_RECLASSIFICADA:
            # ressalva reclassificada como discordancia (ver regra no cabecalho)
            resposta = "Não"

        respostas.append(resposta)
    return respostas


def ler_respostas_do_csv(caminho_csv):
    # O CSV publicado em dados/ tem uma linha por resposta (item x juiz),
    # com a ressalva escrita pelo juiz quando houve. E a mesma informacao
    # da planilha, sem os dados pessoais dos avaliadores. Ver CODEBOOK.md.
    respostas_por_juiz = {}

    with caminho_csv.open(encoding="utf-8", newline="") as arquivo:
        for linha in csv.DictReader(arquivo):
            numero_do_item = int(linha["item"])
            numero_do_juiz = int(linha["juiz"])
            resposta = normalizar_resposta(linha["resposta"])
            tem_ressalva = str(linha.get("tem_ressalva", "")).strip() == "1"

            if resposta == "Sim" and tem_ressalva and numero_do_item in ITENS_COM_RESSALVA_RECLASSIFICADA:
                # ressalva reclassificada como discordancia (ver regra no cabecalho)
                resposta = "Não"

            respostas_por_juiz.setdefault(numero_do_juiz, {})[numero_do_item] = resposta

    if not respostas_por_juiz:
        raise ValueError("o CSV está vazio: " + str(caminho_csv))

    # devolve no mesmo formato da leitura do Excel: uma lista por juiz,
    # com as respostas na ordem dos itens
    return [
        [respostas_por_juiz[juiz][item] for item in sorted(respostas_por_juiz[juiz])]
        for juiz in sorted(respostas_por_juiz)
    ]


def carregar_arquivo(caminho_excel):
    if caminho_excel.suffix.lower() == ".csv":
        return ler_respostas_do_csv(caminho_excel)

    livro = openpyxl.load_workbook(caminho_excel, data_only=True)

    respostas_por_juiz = []
    for nome_da_planilha in livro.sheetnames:
        planilha = livro[nome_da_planilha]
        respostas_por_juiz.append(ler_respostas_de_um_juiz(planilha))

    return respostas_por_juiz


def montar_matriz(respostas_por_juiz):
    # uma linha por item: [quantos Sim, quantos Nao]
    numero_de_itens = len(respostas_por_juiz[0])

    matriz = []
    for i in range(numero_de_itens):
        qtd_sim = 0
        qtd_nao = 0
        for respostas in respostas_por_juiz:
            if respostas[i] == "Sim":
                qtd_sim += 1
            else:
                qtd_nao += 1
        matriz.append([qtd_sim, qtd_nao])

    return matriz


def calcular_pa_observada(matriz, numero_de_juizes):
    # concordancia observada media (Pa) - mesma conta usada dentro do Kappa
    # de Fleiss, so que aqui a gente calcula na mao pra poder reusar no AC1
    r = numero_de_juizes
    soma = 0.0
    for qtd_sim, qtd_nao in matriz:
        soma += (qtd_sim * (qtd_sim - 1) + qtd_nao * (qtd_nao - 1)) / (r * (r - 1))
    return soma / len(matriz)


def calcular_prevalencia_media(matriz, numero_de_juizes):
    # pi = proporcao de "Sim" entre TODAS as respostas dadas (todos os itens,
    # todos os juizes juntos) - e essa media geral que o AC1 usa, diferente
    # do Kappa de Fleiss (que usa a soma dos quadrados das proporcoes)
    total_de_sim = sum(qtd_sim for qtd_sim, qtd_nao in matriz)
    total_de_respostas = len(matriz) * numero_de_juizes
    return total_de_sim / total_de_respostas


def calcular_ac1(matriz, numero_de_juizes):
    pa = calcular_pa_observada(matriz, numero_de_juizes)
    pi = calcular_prevalencia_media(matriz, numero_de_juizes)

    # Pe do AC1 (formula de Gwet pra 2 categorias): 2 * pi * (1 - pi)
    pe = 2 * pi * (1 - pi)

    if pe == 1 or pe == 0:
        # pe = 0 acontece quando pi = 0 ou pi = 1 (todo mundo respondeu a
        # MESMA categoria em TODOS os itens) - nesse caso nao sobra nenhuma
        # discordancia real pra uma estatistica corrigida por acaso medir
        return float("nan"), pa, pe

    ac1 = (pa - pe) / (1 - pe)
    return ac1, pa, pe


def interpretar_forca_concordancia(valor):
    # mesma escala de Landis & Koch (1977) usada no script do kappa - a
    # literatura usa essa mesma regua tambem pra interpretar o AC1
    # (ex.: Wongpakaran et al., 2013)
    if math.isnan(valor):
        return "indefinido"
    if valor < 0:
        return "concordância pior do que o esperado por acaso"
    if valor < 0.20:
        return "concordância leve"
    if valor < 0.40:
        return "concordância razoável"
    if valor < 0.60:
        return "concordância moderada"
    if valor < 0.80:
        return "concordância substancial"
    return "concordância quase perfeita"


def montar_texto_do_relatorio(matriz, numero_de_juizes):
    valor_ac1, pa, pe = calcular_ac1(matriz, numero_de_juizes)

    linhas = []
    linhas.append("=" * 60)
    linhas.append("RESULTADO - AC1 de Gwet - EPMR (juízes especialistas)")
    linhas.append("=" * 60)
    linhas.append("Número de juízes: " + str(numero_de_juizes))
    linhas.append("Número de itens: " + str(len(matriz)))
    linhas.append("")

    linhas.append("Contagem por item (quantos Sim, quantos Não):")
    for i, (qtd_sim, qtd_nao) in enumerate(matriz):
        linhas.append("  item " + str(i + 1) + ": Sim=" + str(qtd_sim) + ", Não=" + str(qtd_nao))

    linhas.append("")
    linhas.append("Concordância observada (Pa): " + str(round(pa, 4)))
    if math.isnan(valor_ac1):
        linhas.append("Concordância esperada ao acaso pelo AC1 (Pe): " + str(round(pe, 4)))
        linhas.append("AC1 de Gwet: NÃO DEU PRA INTERPRETAR (sem nenhuma discordância nos dados)")
    else:
        linhas.append("Concordância esperada ao acaso pelo AC1 (Pe): " + str(round(pe, 4)))
        linhas.append("AC1 de Gwet: " + str(round(valor_ac1, 4)))
        linhas.append("  Interpretação: " + interpretar_forca_concordancia(valor_ac1))
        linhas.append(
            "  (diferente do Kappa de Fleiss, o AC1 não sofre do paradoxo do kappa "
            "quando a prevalência das respostas é desbalanceada - ver Gwet, 2008)"
        )

    linhas.append("=" * 60)

    return "\n".join(linhas)


def salvar_relatorio_em_arquivo(texto_relatorio, agora):
    PASTA_RESULTADOS.mkdir(parents=True, exist_ok=True)

    nome_do_arquivo = "ac1_epmr_" + agora + ".txt"
    caminho_completo = PASTA_RESULTADOS / nome_do_arquivo

    caminho_completo.write_text(texto_relatorio, encoding="utf-8")

    return caminho_completo


def montar_resumo_para_figura(matriz, numero_de_juizes):
    """
    Monta os dados prontos pra virar figura (PNG) e pagina (HTML) - ver
    fleiss_kappa_epmr.py pra explicação completa. Aqui só entra o AC1.
    """
    valor_ac1, _pa, _pe = calcular_ac1(matriz, numero_de_juizes)

    texto_ac1 = "—" if math.isnan(valor_ac1) else str(round(valor_ac1, 3))
    interpretacao = interpretar_forca_concordancia(valor_ac1)

    colunas = ["Itens (N)", "Avaliadores", "AC1 de Gwet", "Interpretação"]
    valores = [
        str(len(matriz)),
        str(numero_de_juizes),
        texto_ac1,
        interpretacao if math.isnan(valor_ac1) else interpretacao.capitalize(),
    ]

    titulo = "AC1 de Gwet: EPMR (juízes especialistas)"

    notas = ["Nota. Interpretação do AC1 segundo a escala de Landis e Koch (1977)."]
    if math.isnan(valor_ac1):
        notas = [
            "AC1 indefinido: sem discordância nos dados pra uma estatística corrigida por "
            "acaso medir (ver EXPLICACAO.md, seção 9)."
        ]

    return {"titulo": titulo, "colunas": colunas, "valores": valores, "notas": notas}


def salvar_imagem_relatorio(resumo, agora):
    """
    Gera uma imagem PNG com uma tabela-resumo do resultado (mesmo estilo
    academico das tabelas do artigo). Ver fleiss_kappa_epmr.py pra
    explicação completa de como essa função funciona.
    """
    PASTA_RESULTADOS.mkdir(parents=True, exist_ok=True)

    plt.rcParams["font.family"] = "serif"

    titulo = resumo["titulo"]
    colunas = resumo["colunas"]
    valores = resumo["valores"]
    notas = resumo["notas"]

    linhas_notas = []
    for nota in notas:
        linhas_notas.extend(textwrap.wrap(nota, width=115))
    texto_notas = "\n".join(linhas_notas)

    altura_titulo = 0.5
    altura_tabela = 1.3
    altura_notas = 0.15 + 0.24 * max(len(linhas_notas), 1)
    altura_total = altura_titulo + altura_tabela + altura_notas

    fig = plt.figure(figsize=(9.0, altura_total))
    fig.patch.set_facecolor("white")

    fig.text(
        0.03,
        1 - (0.32 / altura_total),
        titulo,
        fontsize=11,
        style="italic",
        ha="left",
        va="top",
        wrap=True,
    )

    topo_tabela = 1 - (altura_titulo / altura_total)
    altura_tabela_frac = altura_tabela / altura_total
    ax = fig.add_axes([0.03, topo_tabela - altura_tabela_frac, 0.94, altura_tabela_frac])
    ax.axis("off")

    tabela = ax.table(cellText=[valores], colLabels=colunas, cellLoc="center", loc="center")
    tabela.auto_set_font_size(False)
    tabela.set_fontsize(9.5)
    tabela.auto_set_column_width(col=list(range(len(colunas))))
    tabela.scale(1, 2.2)

    for (linha, _coluna), celula in tabela.get_celld().items():
        celula.set_edgecolor("black")
        if linha == 0:
            celula.set_text_props(weight="bold")
            celula.visible_edges = "BT"
            celula.set_linewidth(1.1)
        else:
            celula.visible_edges = "B"
            celula.set_linewidth(1.1)

    y_notas = topo_tabela - altura_tabela_frac - (0.05 / altura_total)
    fig.text(0.03, y_notas, texto_notas, fontsize=8, ha="left", va="top")

    nome_do_arquivo = "ac1_epmr_" + agora + ".png"
    caminho_completo = PASTA_RESULTADOS / nome_do_arquivo

    fig.savefig(caminho_completo, dpi=200, facecolor="white")
    plt.close(fig)

    return caminho_completo


def salvar_html_relatorio(resumo, nome_da_imagem, agora):
    """
    Gera uma pagina HTML com a mesma tabela-resumo - ver
    fleiss_kappa_epmr.py pra explicação completa. Referencia a imagem
    PNG gerada nessa mesma execução (mesmo "agora"), nunca uma antiga.
    """
    PASTA_RESULTADOS.mkdir(parents=True, exist_ok=True)

    titulo = resumo["titulo"]
    colunas = resumo["colunas"]
    valores = resumo["valores"]
    notas = resumo["notas"]

    celulas_cabecalho = "".join("<th>" + c + "</th>" for c in colunas)
    celulas_dados = "".join("<td>" + v + "</td>" for v in valores)
    paragrafos_notas = "".join("<p>" + n + "</p>" for n in notas)

    html = (
        "<!doctype html>\n"
        '<html lang="pt-BR">\n'
        "<head>\n"
        '<meta charset="utf-8">\n'
        "<title>" + titulo + "</title>\n"
        "<style>\n"
        "  body { margin: 0; padding: 24px; background: #ffffff;\n"
        '    font-family: "Times New Roman", Times, Georgia, serif; color: #000000; }\n'
        "  .titulo-tabela { font-size: 14px; margin-bottom: 14px; font-style: italic; }\n"
        "  table { width: 100%; border-collapse: collapse; font-size: 13px; }\n"
        "  thead tr { border-top: 1.6px solid #000; }\n"
        "  th, td { padding: 7px 10px; text-align: center; }\n"
        "  thead th { border-bottom: 1px solid #000; font-weight: bold; }\n"
        "  tbody tr:last-child td { border-bottom: 1.6px solid #000; }\n"
        "  .nota { font-size: 11.5px; margin-top: 12px; line-height: 1.5; }\n"
        "  .nota p { margin: 4px 0; }\n"
        "  .imagem-recente { margin-top: 28px; }\n"
        "  .imagem-recente p { font-size: 11px; font-style: italic; margin-bottom: 6px; }\n"
        "  .imagem-recente img { max-width: 100%; border: 1px solid #ccc; }\n"
        "</style>\n"
        "</head>\n"
        "<body>\n"
        '  <div class="titulo-tabela">' + titulo + "</div>\n"
        "  <table>\n"
        "    <thead><tr>" + celulas_cabecalho + "</tr></thead>\n"
        "    <tbody><tr>" + celulas_dados + "</tr></tbody>\n"
        "  </table>\n"
        '  <div class="nota">' + paragrafos_notas + "</div>\n"
        '  <div class="imagem-recente">\n'
        "    <p>Imagem gerada nesta execução (" + agora + "):</p>\n"
        '    <img src="' + nome_da_imagem + '" alt="' + titulo + '">\n'
        "  </div>\n"
        "</body>\n"
        "</html>\n"
    )

    nome_do_arquivo = "ac1_epmr_" + agora + ".html"
    caminho_completo = PASTA_RESULTADOS / nome_do_arquivo
    caminho_completo.write_text(html, encoding="utf-8")

    return caminho_completo


def main():
    if len(sys.argv) > 1:
        caminho_excel = Path(sys.argv[1])
    else:
        caminho_excel = ARQUIVO_PADRAO

    if not caminho_excel.exists():
        print("Não achei esse arquivo aqui: " + str(caminho_excel))
        sys.exit(1)

    respostas_por_juiz = carregar_arquivo(caminho_excel)
    matriz = montar_matriz(respostas_por_juiz)

    texto_relatorio = montar_texto_do_relatorio(matriz, len(respostas_por_juiz))

    print(texto_relatorio)

    agora = datetime.now().strftime("%Y%m%d_%H%M%S")

    caminho_salvo = salvar_relatorio_em_arquivo(texto_relatorio, agora)
    print("\nRelatório salvo em: " + str(caminho_salvo))

    resumo = montar_resumo_para_figura(matriz, len(respostas_por_juiz))

    caminho_imagem = salvar_imagem_relatorio(resumo, agora)
    print("Imagem (tabela-resumo pronta pro artigo) salva em: " + str(caminho_imagem))

    caminho_html = salvar_html_relatorio(resumo, caminho_imagem.name, agora)
    print("Página HTML (com a imagem mais recente) salva em: " + str(caminho_html))


if __name__ == "__main__":
    main()
